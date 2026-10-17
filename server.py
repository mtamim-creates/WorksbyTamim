from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer
from email.parser import BytesParser
from email.policy import default
from io import BytesIO
from urllib.parse import parse_qs
from pathlib import Path
import json, re, csv, zipfile, xml.etree.ElementTree as ET
from pypdf import PdfReader, PdfWriter
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import letter
from reportlab.lib.utils import ImageReader
from docx import Document
from openpyxl import load_workbook, Workbook
from PIL import Image
import pymupdf as fitz, pytesseract
from pptx import Presentation
from pptx.util import Inches, Pt
import difflib, tempfile, os, zipfile

ROOT=Path(__file__).parent

def pdf_bytes(text):
    out=BytesIO(); c=canvas.Canvas(out,pagesize=letter); w,h=letter
    for para in text.splitlines() or ['']:
        if h < 55: c.showPage(); h=letter[1]-48
        c.drawString(48,h,para[:125]); h-=15
    c.save(); return out.getvalue()
def page_list(spec,n,preserve_order=False):
    if not spec.strip(): return list(range(n))
    out=[]
    for bit in spec.split(','):
        bit=bit.strip()
        if '-' in bit:
            a,b=bit.split('-',1); out.extend(range(max(1,int(a)),min(n,int(b))+1))
        elif bit: out.append(int(bit))
    valid=[i-1 for i in out if 1<=i<=n]
    return list(dict.fromkeys(valid)) if preserve_order else sorted(set(valid))
def pdf_transform(name,data,opts):
    if name=='PowerPoint to PDF':
        prs=Presentation(BytesIO(data)); lines=[]
        for slide in prs.slides:
            for shape in slide.shapes:
                if getattr(shape,'has_text_frame',False): lines.extend(p.text for p in shape.text_frame.paragraphs)
            lines.append('---')
        return pdf_bytes('\n'.join(lines)),'pdf','application/pdf'
    if name=='JPG to PDF':
        im=Image.open(BytesIO(data)).convert('RGB'); out=BytesIO(); im.save(out,'PDF',resolution=150); return out.getvalue(),'pdf','application/pdf'
    if name=='Scan to PDF':
        im=Image.open(BytesIO(data)).convert('RGB'); out=BytesIO(); im.save(out,'PDF',resolution=150); return out.getvalue(),'pdf','application/pdf'
    if name=='Word to PDF':
        d=Document(BytesIO(data)); return pdf_bytes('\n'.join(p.text for p in d.paragraphs)),'pdf','application/pdf'
    if name=='Excel to PDF':
        wb=load_workbook(BytesIO(data),read_only=True,data_only=True); lines=[]
        for ws in wb.worksheets:
            lines.append(ws.title)
            for row in ws.iter_rows(values_only=True): lines.append('   |   '.join('' if v is None else str(v) for v in row))
        return pdf_bytes('\n'.join(lines)),'pdf','application/pdf'
    if name=='HTML to PDF':
        txt=re.sub(r'<(script|style).*?</\1>',' ',data.decode('utf8','replace'),flags=re.I|re.S); txt=re.sub(r'<[^>]+>',' ',txt); return pdf_bytes(re.sub(r'\s+',' ',txt)),'pdf','application/pdf'
    if name=='PDF to Word':
        reader=PdfReader(BytesIO(data)); doc=Document()
        for i,p in enumerate(reader.pages):
            if i: doc.add_page_break()
            doc.add_paragraph(p.extract_text() or '')
        o=BytesIO();doc.save(o);return o.getvalue(),'docx','application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    if name=='PDF to PowerPoint':
        doc=fitz.open(stream=data,filetype='pdf'); prs=Presentation()
        for page in doc:
            slide=prs.slides.add_slide(prs.slide_layouts[6]); box=slide.shapes.add_textbox(Inches(.6), Inches(.6), Inches(12), Inches(6.3)); tf=box.text_frame; tf.text=page.get_text() or '[No selectable text on this page]'
            for para in tf.paragraphs:
                for run in para.runs: run.font.size=Pt(14)
        o=BytesIO();prs.save(o);return o.getvalue(),'pptx','application/vnd.openxmlformats-officedocument.presentationml.presentation'
    if name=='PDF to JPG':
        doc=fitz.open(stream=data,filetype='pdf'); o=BytesIO()
        with zipfile.ZipFile(o,'w',zipfile.ZIP_DEFLATED) as z:
            for i,p in enumerate(doc,1): z.writestr(f'page-{i:03}.jpg',p.get_pixmap(matrix=fitz.Matrix(1.7,1.7),alpha=False).tobytes('jpg'))
        return o.getvalue(),'zip','application/zip'
    if name=='PDF to Markdown':
        doc=fitz.open(stream=data,filetype='pdf'); md='\n\n---\n\n'.join(p.get_text('text') for p in doc); return md.encode(),'md','text/markdown'
    if name=='AI Summarizer':
        doc=fitz.open(stream=data,filetype='pdf'); text='\n'.join(p.get_text() for p in doc).strip()
        if not text: raise ValueError('No selectable text found. Run OCR PDF first, then summarize the OCR output.')
        sentences=re.split(r'(?<=[.!?])\s+',re.sub(r'\s+',' ',text)); summary='LOCAL EXTRACTIVE SUMMARY\n\n'+'\n'.join('• '+x for x in sentences[:min(8,max(3,len(sentences)//5))])+'\n\nThis extractive summary was generated locally; it is not an LLM analysis.'
        return summary.encode(),'txt','text/plain'
    if name=='OCR PDF':
        doc=fitz.open(stream=data,filetype='pdf'); out=[]
        for i,p in enumerate(doc,1):
            text=p.get_text().strip()
            if not text:
                pix=p.get_pixmap(matrix=fitz.Matrix(2,2),alpha=False); im=Image.open(BytesIO(pix.tobytes('png'))); text=pytesseract.image_to_string(im,lang=opts.get('language','eng'))
            out.append(f'--- Page {i} ---\n{text}')
        return ('\n\n'.join(out)).encode(),'txt','text/plain'
    if name=='Translate PDF':
        raise ValueError('Translation needs a configured translation provider/API key. No third-party text service is configured, so the PDF was not sent anywhere.')
    if name=='Compare PDF':
        raise ValueError('Choose two PDFs to compare. Upload both files in the Compare PDF tool.')
    if name=='PDF to Excel':
        wb=Workbook(); ws=wb.active; ws.title='Extracted text'; reader=PdfReader(BytesIO(data))
        for i,p in enumerate(reader.pages,1): ws.append([f'Page {i}',p.extract_text() or ''])
        o=BytesIO();wb.save(o);return o.getvalue(),'xlsx','application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    if name=='Excel to CSV':
        wb=load_workbook(BytesIO(data),read_only=True,data_only=True); ws=wb.active; o=BytesIO(); import io; t=io.StringIO(); csv.writer(t).writerows(ws.values); return t.getvalue().encode(),'csv','text/csv'
    if name=='CSV to Excel':
        import io; rows=list(csv.reader(io.StringIO(data.decode('utf8-sig')))); wb=Workbook();ws=wb.active
        for r in rows: ws.append(r)
        o=BytesIO();wb.save(o);return o.getvalue(),'xlsx','application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    if name in ('PDF to PDF/A','Compress PDF','Repair PDF','OCR PDF','PDF to JPG','PDF to PowerPoint','PowerPoint to PDF','Translate PDF','AI Summarizer','PDF to Markdown','Compare PDF','Unlock PDF','Protect PDF','Redact PDF','Sign PDF','PDF Forms','Edit PDF','Remove pages','Extract pages','Organize PDF','Split PDF','Merge PDF','Rotate PDF','Add page numbers','Add watermark','Crop PDF'):
        reader=PdfReader(BytesIO(data))
        if reader.is_encrypted:
            pwd=opts.get('password','')
            if not pwd or not reader.decrypt(pwd): raise ValueError('This PDF is password-protected. Enter its password in the options field.')
        writer=PdfWriter(); pages=reader.pages
        if name in ('Split PDF','Extract pages'): pages=[pages[i] for i in page_list(opts.get('pages',''),len(pages))]
        elif name=='Remove pages':
            spec=opts.get('pages','')
            if not spec.strip(): raise ValueError('Enter the page numbers to remove, for example 2, 4-5.')
            remove=set(page_list(spec,len(pages))); pages=[p for i,p in enumerate(pages) if i not in remove]
            if not pages: raise ValueError('Removing those pages would leave an empty PDF.')
        elif name=='Organize PDF':
            order=page_list(opts.get('pages',''),len(pages),preserve_order=True); pages=[pages[i] for i in order] if order else pages
        if name=='Split PDF':
            z=BytesIO()
            with zipfile.ZipFile(z,'w',zipfile.ZIP_DEFLATED) as archive:
                for i,p in enumerate(pages,1):
                    one=PdfWriter();one.add_page(p);b=BytesIO();one.write(b);archive.writestr(f'page-{i:03}.pdf',b.getvalue())
            return z.getvalue(),'zip','application/zip'
        if name=='PDF Forms':
            fields=reader.get_fields() or {}; text='PDF form fields\n\n'+'\n'.join(f'{k}: {v.get("/V", "")}' for k,v in fields.items()) if fields else 'This PDF contains no fillable form fields.'
            return text.encode(),'txt','text/plain'
        if name=='Sign PDF':
            sig=opts.get('text','Signed electronically'); stamp=pdf_bytes(sig); overlay=PdfReader(BytesIO(stamp)).pages[0]
            for p in pages: p.merge_page(overlay)
        if name=='Edit PDF':
            stamp=pdf_bytes(opts.get('text','Edited')); overlay=PdfReader(BytesIO(stamp)).pages[0]
            for p in pages: p.merge_page(overlay)
        if name=='Redact PDF':
            doc=fitz.open(stream=data,filetype='pdf'); page_no=max(1,int(opts.get('page','1')))-1
            if page_no>=len(doc): raise ValueError('Page number is outside this document.')
            page=doc[page_no]; rect=fitz.Rect(float(opts.get('x','40')),float(opts.get('y','40')),float(opts.get('x','40'))+float(opts.get('width','200')),float(opts.get('y','40'))+float(opts.get('height','40'))); page.add_redact_annot(rect,fill=(0,0,0)); page.apply_redactions(); return doc.tobytes(garbage=4,deflate=True),'pdf','application/pdf'
        if name=='Compress PDF':
            doc=fitz.open(stream=data,filetype='pdf'); return doc.tobytes(garbage=4,deflate=True,clean=True),'pdf','application/pdf'
        if name=='PDF to PDF/A': raise ValueError('True archival PDF/A conversion requires a standards validator/conversion engine (not installed); a normal PDF will not be mislabeled as PDF/A.')
        for i,p in enumerate(pages):
            if name=='Rotate PDF': p.rotate(int(opts.get('degrees','90'))%360)
            if name=='Crop PDF':
                try:
                    margin=float(opts.get('margin','24')); p.mediabox.lower_left=(float(p.mediabox.left)+margin,float(p.mediabox.bottom)+margin); p.mediabox.upper_right=(float(p.mediabox.right)-margin,float(p.mediabox.top)-margin)
                except: pass
            if name=='Add watermark':
                stamp=pdf_bytes(opts.get('text','CONFIDENTIAL')); overlay=PdfReader(BytesIO(stamp)).pages[0]; p.merge_page(overlay)
            if name=='Add page numbers':
                stamp=pdf_bytes(f"Page {i+1}"); p.merge_page(PdfReader(BytesIO(stamp)).pages[0])
            writer.add_page(p)
        if name=='Protect PDF': writer.encrypt(opts.get('password','converger'))
        out=BytesIO();writer.write(out);return out.getvalue(),'pdf','application/pdf'
    raise ValueError('This tool needs advanced engine support or an external service and is not available in this local installation yet.')

class Handler(SimpleHTTPRequestHandler):
    def translate_path(self,path):
        if path=='/': path='/index.html'
        return str(ROOT / path.lstrip('/').split('?',1)[0])
    def do_POST(self):
        if self.path!='/api/process': self.send_error(404);return
        raw=self.rfile.read(int(self.headers.get('Content-Length','0')))
        msg=BytesParser(policy=default).parsebytes(b'Content-Type: '+self.headers.get('Content-Type','').encode()+b'\r\nMIME-Version: 1.0\r\n\r\n'+raw)
        fields={}; files=[]
        if msg.is_multipart():
            for part in msg.iter_parts():
                key=part.get_param('name',header='content-disposition'); payload=part.get_payload(decode=True) or b''
                if part.get_filename(): files.append((part.get_filename(),payload))
                elif key: fields[key]=payload.decode('utf8','replace')
        try:
            name=fields.get('tool',''); opts=json.loads(fields.get('options','{}')); 
            if name=='Compare PDF':
                if len(files)<2: raise ValueError('Select two PDFs to compare.')
                a='\n'.join(p.extract_text() or '' for p in PdfReader(BytesIO(files[0][1])).pages).splitlines(); b='\n'.join(p.extract_text() or '' for p in PdfReader(BytesIO(files[1][1])).pages).splitlines(); result='\n'.join(difflib.unified_diff(a,b,fromfile=files[0][0],tofile=files[1][0],lineterm='')).encode(); ext,mime='diff.txt','text/plain'
            elif name=='Merge PDF':
                writer=PdfWriter()
                for fn,data in files: writer.append(PdfReader(BytesIO(data)))
                if len(files)<2: raise ValueError('Select at least two PDF files to merge.')
                o=BytesIO();writer.write(o);result,ext,mime=o.getvalue(),'pdf','application/pdf'
            else:
                if not files: raise ValueError('Select a file first.')
                result,ext,mime=pdf_transform(name,files[0][1],opts)
            outname=re.sub(r'[^\w.-]+','_',Path(files[0][0]).stem if files else name)+'.'+ext
            self.send_response(200);self.send_header('Content-Type',mime);self.send_header('Content-Disposition',f'attachment; filename="{outname}"');self.send_header('Content-Length',str(len(result)));self.send_header('Access-Control-Expose-Headers','Content-Disposition');self.end_headers();self.wfile.write(result)
        except Exception as e:
            b=json.dumps({'error':str(e)}).encode();self.send_response(400);self.send_header('Content-Type','application/json');self.send_header('Content-Length',str(len(b)));self.end_headers();self.wfile.write(b)
    def log_message(self,*args): pass

if __name__=='__main__': ThreadingHTTPServer(('0.0.0.0',8000),Handler).serve_forever()
